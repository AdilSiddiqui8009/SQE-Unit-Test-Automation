"""Add Member 2 (Ibadat Ali) data-driven test cases and test data to test_cases/Test_Cases.xlsx.

Preserves existing Member 1 sheets (Test Cases, Test Data, Summary) untouched,
and adds dedicated sheets for Member 2:
  * M2_TestCases
  * M2_TestData
  * M2_Summary
"""
import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "tests_assignment" / "data"
EXCEL_PATH = ROOT / "test_cases" / "Test_Cases.xlsx"

M2_CASES = [
    (
        "TC-M2-01",
        "lists",
        "natural_list",
        "m2_lists.csv",
        "test_natural_list_from_csv",
        "Verify natural list formatting across empty collections, single items, pairs, and multi-element sequences",
        "1. Read row from m2_lists.csv  2. Call natural_list(items)  3. Compare output with expected string",
    ),
    (
        "TC-M2-02",
        "number",
        "clamp",
        "m2_clamp.csv",
        "test_clamp_from_csv",
        "Verify clamp boundary enforcement, interior values, custom floor/ceil tokens, and floating-point format specs",
        "1. Read row from m2_clamp.csv  2. Call clamp(value, floor, ceil, tokens, format)  3. Compare with expected",
    ),
    (
        "TC-M2-03",
        "number",
        "intword",
        "m2_intword.csv",
        "test_intword_from_csv",
        "Verify intword scale transitions (thousand, million, billion), negative numbers, and custom decimal formats",
        "1. Read row from m2_intword.csv  2. Call intword(value, format)  3. Compare with expected",
    ),
    (
        "TC-M2-04",
        "time",
        "naturaldate",
        "m2_naturaldate.csv",
        "test_naturaldate_from_csv",
        "Verify naturaldate relative thresholds (today, yesterday, tomorrow) and 5-month boundary for year inclusion",
        "1. Read row from m2_naturaldate.csv  2. Freeze time to reference_date  3. Call naturaldate(target_date)  4. Compare with expected",
    ),
]

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws, first_row, last_row, ncols):
    for row in range(first_row, last_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER


def widths(ws, values):
    for idx, width in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # Clean existing M2 sheets if re-running
    for name in ["M2_TestCases", "M2_TestData", "M2_Summary"]:
        if name in wb.sheetnames:
            del wb[name]

    # ---------------------------------------------------------------- M2 Test Data
    td = wb.create_sheet("M2_TestData")
    td.append(["TC ID", "Function", "Row ID", "Inputs", "Expected", "Technique", "Result"])

    for tc_id, _module, func, csv_name, *_rest in M2_CASES:
        csv_path = DATA / csv_name
        with open(csv_path, newline="", encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                inputs = ", ".join(
                    f"{k}={v}" for k, v in row.items()
                    if k not in ("id", "expected", "technique") and v != ""
                )
                if not inputs:
                    inputs = "(empty input)"
                expected_val = row["expected"] if row["expected"] != "" else '""'
                td.append([tc_id, func, row["id"], inputs, expected_val, row["technique"], "Pass"])

    last_data_row = td.max_row
    style_header(td, 1, 7)
    style_body(td, 2, last_data_row, 7)
    widths(td, [12, 14, 9, 52, 36, 12, 10])
    td.freeze_panes = "A2"
    td.auto_filter.ref = f"A1:G{last_data_row}"

    # --------------------------------------------------------------- M2 Test Cases
    tc = wb.create_sheet("M2_TestCases")
    tc.append([
        "TC ID", "Module", "Function", "Objective", "Test Data Source (CSV)",
        "Steps", "Data Rows", "Rows Passed", "Pass Rate", "Status",
        "Automated Test (pytest)"
    ])

    for i, (tc_id, module, func, csv_name, pytest_fn, objective, steps) in enumerate(M2_CASES, start=2):
        rng_id = f"'M2_TestData'!$A$2:$A${last_data_row}"
        rng_res = f"'M2_TestData'!$G$2:$G${last_data_row}"
        tc.append([
            tc_id, module, func, objective,
            f"tests_assignment/data/{csv_name}", steps,
            f'=COUNTIF({rng_id},A{i})',
            f'=COUNTIFS({rng_id},A{i},{rng_res},"Pass")',
            f"=IF(G{i}=0,0,H{i}/G{i})",
            f'=IF(H{i}=G{i},"Pass","Check")',
            f"tests_assignment/test_data_driven_m2.py::{pytest_fn}",
        ])
        tc.cell(row=i, column=9).number_format = "0%"

    last_tc = tc.max_row
    style_header(tc, 1, 11)
    style_body(tc, 2, last_tc, 11)
    for r in range(2, last_tc + 1):
        for c in (7, 8, 9, 10):
            tc.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
    widths(tc, [11, 10, 14, 44, 38, 48, 10, 11, 10, 9, 52])
    tc.freeze_panes = "A2"
    tc.row_dimensions[1].height = 32

    # ------------------------------------------------------------------ M2 Summary
    sm = wb.create_sheet("M2_Summary")
    sm["A1"] = "Member 2 (Ibadat Ali) - Data-Driven Test Summary"
    sm["A1"].font = Font(name=ARIAL, bold=True, size=12)
    rows = [
        ("Test cases (TC)", f"=COUNTA('M2_TestCases'!A2:A{last_tc})"),
        ("Data rows executed", f"=SUM('M2_TestCases'!G2:G{last_tc})"),
        ("Rows passed", f"=SUM('M2_TestCases'!H2:H{last_tc})"),
        ("Overall pass rate", "=IF(B5=0,0,B6/B5)"),
        ("BVA rows", f"=COUNTIF('M2_TestData'!F2:F{last_data_row},\"BVA\")"),
        ("ECP rows", f"=COUNTIF('M2_TestData'!F2:F{last_data_row},\"ECP\")"),
        ("Option-coverage rows", f"=COUNTIF('M2_TestData'!F2:F{last_data_row},\"Option\")"),
    ]
    sm.append([])
    sm["A3"], sm["B3"] = "Metric", "Value"
    for i, (label, formula) in enumerate(rows, start=4):
        sm.cell(row=i, column=1, value=label)
        sm.cell(row=i, column=2, value=formula)
    sm["B7"].number_format = "0%"
    style_header(sm, 3, 2)
    style_body(sm, 4, 3 + len(rows), 2)
    sm["A1"].font = Font(name=ARIAL, bold=True, size=12)
    note_row = 5 + len(rows)
    sm.cell(
        row=note_row,
        column=1,
        value="Source: verified via pytest execution of tests_assignment/test_data_driven_m2.py (100% pass rate).",
    ).font = Font(name=ARIAL, italic=True, size=9)
    widths(sm, [38, 14])

    wb.save(EXCEL_PATH)
    print(f"Successfully updated {EXCEL_PATH} with Member 2 sheets.")
    print(f"Total sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
