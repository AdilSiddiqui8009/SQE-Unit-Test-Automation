"""Build test_cases/Test_Cases.xlsx from the CSV data files and a pytest JUnit run.

Usage (from repo root):
    pytest tests_assignment/test_data_driven_csv.py --junitxml=reports/data_driven.xml -q
    python scripts/build_excel.py
"""
import csv
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "tests_assignment" / "data"
JUNIT = ROOT / "reports" / "data_driven.xml"
OUT = ROOT / "test_cases" / "Test_Cases.xlsx"

# TC id, module, function, csv file, pytest function, objective, steps
CASES = [
    ("TC-DD-001", "number", "ordinal", "ordinal_cases.csv", "test_ordinal_from_csv",
     "Verify ordinal suffix rules (st/nd/rd/th, teen exceptions) and invalid-input handling",
     "1. Read a row from ordinal_cases.csv  2. Call ordinal(value)  3. Compare with expected"),
    ("TC-DD-002", "number", "intcomma", "intcomma_cases.csv", "test_intcomma_from_csv",
     "Verify thousands separators, negatives, floats, ndigits rounding and invalid input",
     "1. Read a row from intcomma_cases.csv  2. Call intcomma(value[, ndigits])  3. Compare with expected"),
    ("TC-DD-003", "number", "apnumber", "apnumber_cases.csv", "test_apnumber_from_csv",
     "Verify AP-style spelling of 0-9 and the 9/10 and 0/-1 boundaries",
     "1. Read a row from apnumber_cases.csv  2. Call apnumber(value)  3. Compare with expected"),
    ("TC-DD-004", "filesize", "naturalsize", "naturalsize_cases.csv", "test_naturalsize_from_csv",
     "Verify decimal/binary/GNU size formatting, unit boundaries and rounding carry-over",
     "1. Read a row from naturalsize_cases.csv  2. Call naturalsize(value, binary, gnu, format)  3. Compare with expected"),
    ("TC-DD-005", "time", "naturaldelta", "naturaldelta_cases.csv", "test_naturaldelta_from_csv",
     "Verify natural duration text at every unit boundary, plus months and minimum_unit options",
     "1. Read a row from naturaldelta_cases.csv  2. Call naturaldelta(timedelta(seconds), months, minimum_unit)  3. Compare with expected"),
    ("TC-DD-006", "time", "precisedelta", "precisedelta_cases.csv", "test_precisedelta_from_csv",
     "Verify exact multi-unit duration text, unit rollovers, minimum_unit and format options",
     "1. Read a row from precisedelta_cases.csv  2. Call precisedelta(timedelta(seconds), minimum_unit, format)  3. Compare with expected"),
]

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def junit_results():
    """Map CSV row id (e.g. 'O-01') -> 'Pass'/'Fail'/'Skip' from the JUnit file."""
    results = {}
    if not JUNIT.exists():
        return results
    for case in ET.parse(JUNIT).getroot().iter("testcase"):
        name = case.get("name", "")
        if "[" not in name:
            continue
        row_id = name[name.index("[") + 1 : -1]
        if case.find("failure") is not None or case.find("error") is not None:
            results[row_id] = "Fail"
        elif case.find("skipped") is not None:
            results[row_id] = "Skip"
        else:
            results[row_id] = "Pass"
    return results


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws, first_row, last_row, ncols):
    for row in range(first_row, last_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER


def widths(ws, values):
    for idx, width in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def main():
    results = junit_results()
    wb = Workbook()

    # ------------------------------------------------------------------ Test Data
    td = wb.active
    td.title = "Test Data"
    td.append(["TC ID", "Function", "Row ID", "Inputs", "Expected", "Technique", "Result"])
    for tc_id, _module, func, csv_name, *_rest in CASES:
        with open(DATA / csv_name, newline="", encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                inputs = ", ".join(
                    f"{k}={v}" for k, v in row.items()
                    if k not in ("id", "expected", "technique") and v != ""
                )
                td.append([tc_id, func, row["id"], inputs, row["expected"],
                           row["technique"], results.get(row["id"], "Not run")])
    last_data_row = td.max_row
    style_header(td, 1, 7)
    style_body(td, 2, last_data_row, 7)
    widths(td, [12, 14, 9, 52, 36, 12, 10])
    td.freeze_panes = "A2"
    td.auto_filter.ref = f"A1:G{last_data_row}"

    # ----------------------------------------------------------------- Test Cases
    tc = wb.create_sheet("Test Cases", 0)
    tc.append(["TC ID", "Module", "Function", "Objective", "Test Data Source (CSV)",
               "Steps", "Data Rows", "Rows Passed", "Pass Rate", "Status",
               "Automated Test (pytest)"])
    for i, (tc_id, module, func, csv_name, pytest_fn, objective, steps) in enumerate(CASES, start=2):
        rng_id = f"'Test Data'!$A$2:$A${last_data_row}"
        rng_res = f"'Test Data'!$G$2:$G${last_data_row}"
        tc.append([
            tc_id, module, func, objective,
            f"tests_assignment/data/{csv_name}", steps,
            f'=COUNTIF({rng_id},A{i})',
            f'=COUNTIFS({rng_id},A{i},{rng_res},"Pass")',
            f"=IF(G{i}=0,0,H{i}/G{i})",
            f'=IF(H{i}=G{i},"Pass","Check")',
            f"tests_assignment/test_data_driven_csv.py::{pytest_fn}",
        ])
        tc.cell(row=i, column=9).number_format = "0%"
    last_tc = tc.max_row
    style_header(tc, 1, 11)
    style_body(tc, 2, last_tc, 11)
    for r in range(2, last_tc + 1):
        for c in (7, 8, 9, 10):
            tc.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
    widths(tc, [11, 10, 14, 44, 38, 48, 10, 11, 10, 9, 48])
    tc.freeze_panes = "A2"
    tc.row_dimensions[1].height = 32

    # -------------------------------------------------------------------- Summary
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Data-driven test summary (all figures are formulas over 'Test Data')"
    sm["A1"].font = Font(name=ARIAL, bold=True, size=12)
    rows = [
        ("Test cases (TC)", f"=COUNTA('Test Cases'!A2:A{last_tc})"),
        ("Data rows executed", f"=SUM('Test Cases'!G2:G{last_tc})"),
        ("Rows passed", f"=SUM('Test Cases'!H2:H{last_tc})"),
        ("Overall pass rate", "=IF(B5=0,0,B6/B5)"),
        ("BVA rows", f"=COUNTIF('Test Data'!F2:F{last_data_row},\"BVA\")"),
        ("ECP rows", f"=COUNTIF('Test Data'!F2:F{last_data_row},\"ECP\")"),
        ("Option-coverage rows", f"=COUNTIF('Test Data'!F2:F{last_data_row},\"Option\")"),
        ("Negative / error-guessing rows", f"=COUNTIF('Test Data'!F2:F{last_data_row},\"Negative\")"),
    ]
    sm.append([])
    sm["A3"], sm["B3"] = "Metric", "Value"
    for i, (label, formula) in enumerate(rows, start=4):
        sm.cell(row=i, column=1, value=label)
        sm.cell(row=i, column=2, value=formula)
    sm["B7"].number_format = "0%"
    style_header(sm, 3, 2)
    style_body(sm, 4, 3 + len(rows), 2)
    sm["A1"].font = Font(name=ARIAL, bold=True, size=12)
    note_row = 5 + len(rows)
    sm.cell(row=note_row, column=1,
            value="Source: results come from pytest JUnit output (reports/data_driven.xml); "
                  "regenerate with scripts/build_excel.py.").font = Font(name=ARIAL, italic=True, size=9)
    widths(sm, [38, 14])

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({last_data_row - 1} data rows, {len(CASES)} test cases)")


if __name__ == "__main__":
    main()
